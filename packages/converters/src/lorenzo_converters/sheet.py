"""A spreadsheet inventory into a ledger (ADR 0243).

The sheet is an Items table (name, count, weight of one, value, container, kind, description) and
optionally a Container table (name, where it is kept), as an `.xlsx`, or the Items table alone as a
`.csv`. Columns are found by their headers, in German or English. Nothing is guessed: a cell that is
empty or a formula is left out of its line, and what is odd is said as a warning.
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any

from lorenzo_converters.ledger import (
    MAX_DEPTH,
    MAX_LINES,
    MAX_QUANTITY,
    Ledger,
    Line,
    count,
    depth,
)

# What a header may be called, compared without case, surrounding space or a closing dot or colon.
HEADERS: dict[str, tuple[str, ...]] = {
    "name": ("name", "gegenstand", "item", "bezeichnung", "artikel"),
    "quantity": ("anzahl", "count", "qty", "quantity", "menge", "stück", "stueck", "#"),
    "weight": ("gewicht", "weight"),
    "value": ("gesch. wert", "geschätzter wert", "wert", "value", "preis", "price"),
    "container": ("container", "behälter", "behaelter"),
    "kind": ("art", "kind", "type", "typ", "kategorie", "category"),
    "note": ("beschreibung", "description", "note", "notes", "notiz", "bemerkung", "anmerkung"),
    "place": ("ort", "place", "location", "standort"),
}
ITEM_COLUMNS = ("quantity", "weight", "value", "container", "kind", "note")
# What a container may be called to mean "in hand or on the character" (the format's own names).
EQUIPPED = ("holding", "equipped", "hands", "in hand", "worn")
# A table in Excel can run to a million rows of nothing: stop after this many blank names.
BLANK_RUN = 200
# A stack at the top level is written as this many lines at most; more is left for the GM to box.
UNSTACK_LIMIT = 10

_POUNDS = re.compile(r"^(\d+(?:[.,]\d+)?)\s*(?:lbs?\.?|pounds?)?$", re.IGNORECASE)


class ConversionError(Exception):
    """The file cannot be converted; the message says why."""


@dataclass
class Result:
    ledger: Ledger
    warnings: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


@dataclass
class Row:
    number: int
    name: str
    quantity: int = 1
    weight: float | None = None
    value: str | None = None
    kind: str | None = None
    note: str | None = None
    container: str | None = None


def _key(text: str) -> str:
    return " ".join(text.split()).casefold()


def _header(text: object) -> str:
    return _key(str(text)).rstrip(".:") if text is not None else ""


def _columns(header: Iterable[object]) -> dict[str, int]:
    found: dict[str, int] = {}
    for at, cell in enumerate(header):
        label = _header(cell)
        for column, names in HEADERS.items():
            if column not in found and (label in names or label.rstrip(".") in names):
                found[column] = at
    return found


# ---- reading --------------------------------------------------------------------------------


@dataclass
class Table:
    columns: dict[str, int]
    rows: list[tuple[int, tuple[Any, ...]]]


def _tables(rows: Iterator[tuple[Any, ...]]) -> Table | None:
    """The first row with two or more text cells is the header; the rows after it are the table."""
    header: tuple[Any, ...] | None = None
    number = 0
    body: list[tuple[int, tuple[Any, ...]]] = []
    blank = 0
    columns: dict[str, int] = {}
    for row in rows:
        number += 1
        if header is None:
            if sum(isinstance(cell, str) and cell.strip() != "" for cell in row) >= 2:
                header = row
                columns = _columns(row)
            continue
        name = row[columns["name"]] if "name" in columns and columns["name"] < len(row) else None
        if name is None or str(name).strip() == "":
            blank += 1
            if blank >= BLANK_RUN:
                break
            continue
        blank = 0
        body.append((number, row))
    return Table(columns, body) if header is not None else None


def _read_xlsx(path: Path) -> tuple[Table | None, Table | None, list[str]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(path, read_only=True, data_only=False)
    except Exception as error:  # a damaged or foreign file, whatever openpyxl calls it
        raise ConversionError(f"{path.name} cannot be read as an Excel file: {error}") from error
    items: Table | None = None
    places: Table | None = None
    seen: list[str] = []
    try:
        for sheet in book.worksheets:
            table = _tables(sheet.iter_rows(values_only=True))
            if table is None or "name" not in table.columns:
                continue
            seen.append(f"{sheet.title}: {', '.join(sorted(table.columns))}")
            if any(column in table.columns for column in ITEM_COLUMNS):
                if items is None:
                    items = table
            elif "place" in table.columns and places is None:
                places = table
    finally:
        book.close()
    return items, places, seen


def _read_csv(path: Path) -> tuple[Table | None, list[str]]:
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    table = _tables(iter(tuple(row) for row in csv.reader(text.splitlines(), dialect)))
    seen = [f"columns: {', '.join(sorted(table.columns))}"] if table else []
    return table, seen


# ---- cells ----------------------------------------------------------------------------------


class _Cells:
    def __init__(self) -> None:
        self.warnings: list[str] = []

    def _text(self, cell: object) -> str | None:
        if cell is None:
            return None
        if isinstance(cell, float) and cell.is_integer():
            cell = int(cell)
        text = str(cell).strip()
        return text or None

    def text(self, cell: object, where: str, what: str) -> str | None:
        text = self._text(cell)
        if text is not None and text.startswith("="):
            self.warnings.append(f"{where}: {what} is a formula, left out")
            return None
        return text

    def quantity(self, cell: object, where: str) -> int:
        text = self.text(cell, where, "the count")
        if text is None:
            return 1
        try:
            number = float(text.replace(",", "."))
        except ValueError:
            self.warnings.append(f"{where}: the count {text!r} is not a number, taken as 1")
            return 1
        if number < 1 or not number.is_integer():
            self.warnings.append(f"{where}: the count {text!r} is not a whole number, taken as 1")
            return 1
        return int(number)

    def weight(self, cell: object, where: str) -> float | None:
        text = self.text(cell, where, "the weight")
        if text is None:
            return None
        match = _POUNDS.match(text)
        if not match:
            self.warnings.append(
                f"{where}: the weight {text!r} is not a number of pounds, left out"
            )
            return None
        return float(match.group(1).replace(",", "."))


def _rows(table: Table) -> tuple[list[Row], list[str]]:
    cells = _Cells()
    columns = table.columns
    rows: list[Row] = []
    for number, raw in table.rows:

        def cell(column: str, raw: tuple[Any, ...] = raw) -> object:
            at = columns.get(column)
            return raw[at] if at is not None and at < len(raw) else None

        where = f"row {number}"
        name = cells.text(cell("name"), where, "the name")
        if name is None:
            continue
        quantity = cells.quantity(cell("quantity"), f"row {number} ({name})")
        rows.append(
            Row(
                number=number,
                name=name,
                quantity=quantity,
                weight=cells.weight(cell("weight"), f"row {number} ({name})"),
                value=cells.text(cell("value"), f"row {number} ({name})", "the value"),
                kind=cells.text(cell("kind"), f"row {number} ({name})", "the kind"),
                note=cells.text(cell("note"), f"row {number} ({name})", "the description"),
                container=cells.text(cell("container"), f"row {number} ({name})", "the container"),
            )
        )
    return rows, cells.warnings


# ---- building -------------------------------------------------------------------------------


@dataclass
class _Node:
    line: Line
    parent: str | None = None  # the key of the container it is in
    claimed: bool = True


def _line(row: Row) -> Line:
    line = Line(
        name=row.name,
        quantity=row.quantity,
        weight=row.weight,
        value=row.value,
        kind=row.kind,
    )
    if row.note:
        if "\n" in row.note.strip():
            line.description = "\n".join(part.rstrip() for part in row.note.strip().split("\n"))
        else:
            line.note = row.note
    return line


def convert(path: Path, owner: str, home: str | None = None, equipped: str | None = None) -> Result:
    """The ledger for the sheet at `path`, and what was said about it on the way."""
    suffix = path.suffix.lower()
    places_table: Table | None = None
    if suffix == ".xlsx":
        items_table, places_table, seen = _read_xlsx(path)
    elif suffix in (".csv", ".tsv", ".txt"):
        items_table, seen = _read_csv(path)
    else:
        raise ConversionError(f"{path.name}: expected an .xlsx or .csv file")
    if items_table is None or "name" not in items_table.columns:
        found = "; ".join(seen) or "no table with a header row"
        raise ConversionError(
            f"{path.name}: no table of items found (it needs a name column and at least one of "
            f"count, weight, value, container, kind, description). Found: {found}"
        )

    rows, warnings = _rows(items_table)
    notes: list[str] = []
    equipped_keys = {*EQUIPPED, *([_key(equipped)] if equipped else [])}

    # Where containers are kept, by key, and the names the Container table gives them.
    place_of: dict[str, str] = {}
    sheet_names: dict[str, str] = {}
    if places_table is not None:
        for _number, raw in places_table.rows:
            name = raw[places_table.columns["name"]]
            at = places_table.columns["place"]
            place = raw[at] if at < len(raw) else None
            key = _key(str(name))
            if key in equipped_keys:
                continue
            sheet_names.setdefault(key, str(name).strip())
            if place is not None and str(place).strip():
                place_of[key] = str(place).strip()

    home_key: str | None = None
    if home:
        home_key = _key(home)
    elif place_of:
        common = Counter(_key(place) for place in place_of.values()).most_common(1)[0][0]
        home_key = common
        shown = next(place for place in place_of.values() if _key(place) == common)
        notes.append(f"{shown!r} is taken as where the character is; pass --home to change it")

    # A name is a container if the Container table lists it or a row says something is in it.
    referenced: dict[str, str] = {}
    for row in rows:
        if row.container and _key(row.container) not in equipped_keys:
            referenced.setdefault(_key(row.container), row.container)
    container_keys = [*sheet_names, *(k for k in referenced if k not in sheet_names)]

    nodes: dict[str, _Node] = {}
    ordered: list[_Node] = []
    for row in rows:
        key = _key(row.name)
        parent = _key(row.container) if row.container else None
        node = _Node(_line(row), parent)
        if key in container_keys:
            if key in nodes:
                warnings.append(
                    f"row {row.number} ({row.name}): there is already a container of that name, "
                    "this one is kept as an item"
                )
            else:
                nodes[key] = node
        ordered.append(node)
    for key in container_keys:
        if key not in nodes:
            name = sheet_names.get(key) or referenced.get(key, key)
            nodes[key] = _Node(Line(name=name), None, claimed=False)
            ordered.append(nodes[key])
            if key not in sheet_names:
                warnings.append(f"{name!r} is named as a container but is not a row or listed")

    # A container kept somewhere other than home is at the top of Not carried, with its place.
    for key, node in nodes.items():
        place = place_of.get(key)
        if place is not None and home_key is not None and _key(place) != home_key:
            node.line.place = place
            if node.parent is not None:
                warnings.append(
                    f"{node.line.name!r} is in {node.parent!r} but its place is {place!r}: "
                    "kept at the top of Not carried, with its place"
                )
                node.parent = None

    # A loop of containers is cut where it closes.
    for node in ordered:
        seen_keys: set[int] = set()
        current: _Node | None = node
        while current is not None and current.parent is not None:
            if id(current) in seen_keys:
                warnings.append(
                    f"{current.line.name!r} is inside itself through its containers: "
                    "kept at the top level"
                )
                current.parent = None
                break
            seen_keys.add(id(current))
            current = nodes.get(current.parent)

    ledger = Ledger(owner=owner)
    for node in ordered:
        if node.parent is None:
            ledger.not_carried.append(node.line)
        elif node.parent in equipped_keys:
            ledger.equipped.append(node.line)
        else:
            holder = nodes.get(node.parent)
            if holder is None:
                ledger.not_carried.append(node.line)
            else:
                holder.line.contents.append(node.line)

    # The format wants a stack inside something, and a container written once (ADR 0191).
    for top in (ledger.equipped, ledger.not_carried):
        at = 0
        while at < len(top):
            line = top[at]
            if line.contents and line.quantity != 1:
                warnings.append(
                    f"{line.name!r}: {line.quantity} of a thing that holds things; the importer "
                    "wants it written once for each"
                )
            elif line.quantity > 1 and not line.contents:
                if line.quantity <= UNSTACK_LIMIT:
                    warnings.append(
                        f"{line.quantity} x {line.name!r} is at the top level, where a stack "
                        f"is not allowed: written as {line.quantity} lines"
                    )
                    top[at : at + 1] = [replace(line, quantity=1) for _ in range(line.quantity)]
                    at += line.quantity
                    continue
                warnings.append(
                    f"{line.quantity} x {line.name!r} is at the top level, where a stack is not "
                    "allowed: put it in a container"
                )
            at += 1

    total = count(ledger.equipped) + count(ledger.not_carried)
    if total > MAX_LINES:
        warnings.append(f"{total} lines: the importer takes at most {MAX_LINES}")
    deepest = max(depth(ledger.equipped), depth(ledger.not_carried))
    if deepest > MAX_DEPTH:
        warnings.append(f"containers {deepest} deep: the importer takes at most {MAX_DEPTH}")
    for row in rows:
        if row.quantity > MAX_QUANTITY:
            warnings.append(
                f"row {row.number} ({row.name}): {row.quantity} of one thing, "
                f"the importer takes at most {MAX_QUANTITY}"
            )
    notes.append(f"{total} lines for {owner}")
    return Result(ledger, warnings, notes)
